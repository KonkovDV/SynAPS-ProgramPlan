from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from synaps_programplan.io.excel import read_excel, write_excel, write_template
from synaps_programplan.model import DependencyType, RiskDriver, TaskKind, TaskStatus
from tests.conftest import dep, needs, person, program, stand, task, uses


def test_excel_roundtrip(tmp_path: Path) -> None:
    original = program(
        [
            task("a", 3, demands=uses("st") + needs("design")),
            task("m", 0),
        ],
        [dep("a", "m", "FS", 1)],
        resources=[stand(), person("ivanov", ["design"])],
        skills=["design"],
    )
    path = tmp_path / "program.xlsx"
    write_excel(original, path)
    loaded = read_excel(path)
    assert [(t.id, t.duration_wd, t.project_id) for t in loaded.tasks] == [
        ("a", 3, "p1"),
        ("m", 0, "p1"),
    ]
    assert loaded.tasks[0].demands[0].resource_id == "st"
    assert loaded.tasks[0].demands[1].skill_id == "design"
    assert loaded.dependencies[0].lag_wd == 1
    assert {r.id for r in loaded.resources} == {"st", "ivanov"}
    assert loaded.resources[1].skills == ["design"]


def test_risk_register_roundtrip(tmp_path: Path) -> None:
    original = program(
        [task("a", 3), task("b", 2)],
        risk_drivers=[RiskDriver(id="R1", name="Повторные испытания", probability=0.3, task_ids=["a", "b"])],
    )
    path = tmp_path / "program.xlsx"
    write_excel(original, path)
    book = load_workbook(path)
    book["Risks"].append(["R2", "Поставка", "0,25", 1, None, 2, "b", "снабжение"])
    book.save(path)
    loaded = read_excel(path)
    first, second = loaded.risk_drivers
    assert first.task_ids == ["a", "b"] and first.probability == 0.3
    assert second.probability == 0.25 and second.mode == 1.2 and second.owner == "снабжение"


def test_optional_columns_may_stay_empty(tmp_path: Path) -> None:
    path = tmp_path / "manual.xlsx"
    write_template(path)
    book = load_workbook(path)
    book["Program"].append(["prog", "Программа", "2027-01-11", "2027-12-30", None])
    book["Projects"].append(["p1", None, "ОКР-1"])
    book["Tasks"].append(["a", "p1", None, "Расчёт", 5])
    book["Tasks"].append(["m", "p1", None, "Веха", 0])
    book["Dependencies"].append(["a", "m"])
    book["Resources"].append(["st", "STAND", None, "Стенд", 1])
    book["Demands"].append(["a", "st", None, None])
    book.save(path)
    loaded = read_excel(path)
    assert loaded.projects[0].code == "p1"
    assert [t.kind for t in loaded.tasks] == [TaskKind.WORK, TaskKind.MILESTONE]
    assert loaded.tasks[0].status is TaskStatus.PLANNED
    assert loaded.dependencies[0].type is DependencyType.FS
    assert loaded.tasks[0].demands[0].units == 1


def test_template_has_no_data_rows(tmp_path: Path) -> None:
    path = tmp_path / "template.xlsx"
    write_template(path)
    try:
        read_excel(path)
    except ValueError as exc:
        assert "Program" in str(exc)
    else:
        raise AssertionError("comment row was imported as data")
