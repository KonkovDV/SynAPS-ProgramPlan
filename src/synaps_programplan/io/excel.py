"""Excel workbook for one consolidated program (plan section T1.3).

Sheets: Program, Projects, WBS, Tasks, Demands, Dependencies, Resources, Skills,
Exceptions, Risks (optional). Dates are ISO ``YYYY-MM-DD``. Several skills in one cell are
separated by ``;``. An empty id cell or a row whose first cell starts with
``#`` is skipped, so the template can carry a comment row.
"""

from __future__ import annotations

import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from synaps_programplan.model import (
    Calendar,
    CalendarBase,
    CapacityException,
    Demand,
    Dependency,
    DependencySource,
    DependencyType,
    ExceptionReason,
    OKRProgram,
    Program,
    Project,
    Provenance,
    ProvenanceKind,
    Resource,
    ResourceKind,
    RiskDriver,
    Skill,
    Task,
    TaskKind,
    TaskStatus,
    WBSKind,
    WBSNode,
)

MAX_XLSX_FILES = 64
MAX_XLSX_UNCOMPRESSED = 32 * 1024 * 1024
MAX_XLSX_RATIO = 100


def _refuse_expanding_workbook(path: Path) -> None:
    """Stop a zip bomb before openpyxl unpacks it."""
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise ValueError(f"{path.name} is not an Excel workbook") from exc
    with archive:
        parts = archive.infolist()
        if len(parts) > MAX_XLSX_FILES:
            raise ValueError(f"{path.name} contains too many parts to be read")
        expanded = sum(part.file_size for part in parts)
        packed = sum(part.compress_size for part in parts) or 1
        if expanded > MAX_XLSX_UNCOMPRESSED or expanded / packed > MAX_XLSX_RATIO:
            raise ValueError(f"{path.name} expands too far to be read")


_SHEETS: dict[str, list[str]] = {
    "Program": ["id", "name", "horizon_start", "horizon_end", "status_date"],
    "Projects": ["id", "code", "name", "priority", "due_date", "deadline", "enterprise"],
    "WBS": ["id", "project_id", "parent_id", "code", "name", "kind"],
    "Tasks": [
        "id",
        "project_id",
        "wbs_id",
        "name",
        "duration_wd",
        "kind",
        "status",
        "earliest_start",
        "latest_finish",
        "due_date",
        "deadline",
        "shift_limit_wd",
        "pinned",
        "remaining_wd",
        "actual_start",
        "actual_finish",
        "planned_start",
        "planned_finish",
        "okr_stage",
    ],
    "Demands": ["task_id", "resource_id", "skill_id", "units"],
    "Dependencies": ["src_task_id", "dst_task_id", "type", "lag_wd", "max_lag_wd", "hard", "source"],
    "Resources": ["id", "kind", "code", "name", "capacity_units", "skills", "org_unit"],
    "Skills": ["id", "code", "name"],
    "Exceptions": ["resource_id", "start", "end", "units_available", "reason"],
    "Risks": ["id", "name", "probability", "low", "mode", "high", "task_ids", "owner", "group"],
}


def write_template(path: Path) -> None:
    """Header row plus one commented example, so a planner sees every column."""
    book = Workbook()
    first = True
    for title, headers in _SHEETS.items():
        sheet = book.active if first else book.create_sheet(title)
        first = False
        sheet.title = title
        sheet.append(headers)
        sheet.append([f"# {headers[0]}"] + [None] * (len(headers) - 1))
        sheet.freeze_panes = "A2"
    book.save(path)


def write_excel(program: OKRProgram, path: Path) -> None:
    book = Workbook()
    book.remove(book.active)
    rows = _rows(program)
    for title, headers in _SHEETS.items():
        sheet = book.create_sheet(title)
        sheet.append(headers)
        for row in rows[title]:
            sheet.append([row.get(header) for header in headers])
        sheet.freeze_panes = "A2"
    book.save(path)


def read_excel(path: Path, *, provenance: Provenance | None = None) -> OKRProgram:
    _refuse_expanding_workbook(path)
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        tables = {name: _table(book, name) for name in _SHEETS}
    finally:
        book.close()
    if not tables["Program"]:
        raise ValueError(f"{path.name}: sheet Program has no data row")
    header = tables["Program"][0]
    demands: dict[str, list[Demand]] = {}
    for row in tables["Demands"]:
        demands.setdefault(_text(row, "task_id"), []).append(
            Demand(
                resource_id=_opt(row, "resource_id"),
                skill_id=_opt(row, "skill_id"),
                units=_int(row, "units", default=1),
            )
        )
    tasks = []
    for row in tables["Tasks"]:
        duration = _int(row, "duration_wd")
        kind = TaskKind(_opt(row, "kind") or ("MILESTONE" if duration == 0 else "WORK"))
        tasks.append(
            Task(
                id=_text(row, "id"),
                project_id=_text(row, "project_id"),
                wbs_id=_opt(row, "wbs_id"),
                name=_text(row, "name"),
                duration_wd=duration,
                kind=kind,
                demands=demands.get(_text(row, "id"), []),
                earliest_start=_date(row, "earliest_start"),
                latest_finish=_date(row, "latest_finish"),
                due_date=_date(row, "due_date"),
                deadline=_date(row, "deadline"),
                shift_limit_wd=_optional_int(row, "shift_limit_wd"),
                pinned=_bool(row, "pinned"),
                status=TaskStatus(_opt(row, "status") or "PLANNED"),
                remaining_wd=_optional_int(row, "remaining_wd"),
                actual_start=_date(row, "actual_start"),
                actual_finish=_date(row, "actual_finish"),
                planned_start=_date(row, "planned_start"),
                planned_finish=_date(row, "planned_finish"),
                okr_stage=_opt(row, "okr_stage"),
            )
        )
    return OKRProgram(
        program=Program(
            id=_text(header, "id"),
            name=_text(header, "name"),
            calendar_id="ru",
            horizon_start=_date(header, "horizon_start") or date.today(),
            horizon_end=_date(header, "horizon_end") or date.today(),
            status_date=_date(header, "status_date"),
        ),
        calendars=[Calendar(id="ru", base=CalendarBase.RU_PRODUCTION)],
        projects=[
            Project(
                id=_text(row, "id"),
                code=_opt(row, "code") or _text(row, "id"),
                name=_text(row, "name"),
                priority=_int(row, "priority", default=500),
                due_date=_date(row, "due_date"),
                deadline=_date(row, "deadline"),
                enterprise=_opt(row, "enterprise"),
            )
            for row in tables["Projects"]
        ],
        wbs=[
            WBSNode(
                id=_text(row, "id"),
                project_id=_text(row, "project_id"),
                parent_id=_opt(row, "parent_id"),
                code=_opt(row, "code") or _text(row, "id"),
                name=_text(row, "name"),
                kind=WBSKind(_opt(row, "kind") or "PACKAGE"),
            )
            for row in tables["WBS"]
        ],
        tasks=tasks,
        dependencies=[
            Dependency(
                src_task_id=_text(row, "src_task_id"),
                dst_task_id=_text(row, "dst_task_id"),
                type=DependencyType(_opt(row, "type") or "FS"),
                lag_wd=_int(row, "lag_wd", default=0),
                max_lag_wd=_optional_int(row, "max_lag_wd"),
                hard=_bool(row, "hard", default=True),
                source=DependencySource(_opt(row, "source") or "IMPORTED"),
            )
            for row in tables["Dependencies"]
        ],
        resources=[
            Resource(
                id=_text(row, "id"),
                kind=ResourceKind(_text(row, "kind")),
                code=_opt(row, "code") or _text(row, "id"),
                name=_text(row, "name"),
                capacity_units=_int(row, "capacity_units", default=1),
                skills=[part for part in (_opt(row, "skills") or "").split(";") if part],
                org_unit=_opt(row, "org_unit"),
            )
            for row in tables["Resources"]
        ],
        skills=[
            Skill(id=_text(row, "id"), code=_opt(row, "code") or _text(row, "id"), name=_text(row, "name"))
            for row in tables["Skills"]
        ],
        capacity_exceptions=[
            CapacityException(
                resource_id=_text(row, "resource_id"),
                start=_date(row, "start") or date.today(),
                end=_date(row, "end") or date.today(),
                units_available=_int(row, "units_available", default=0),
                reason=ExceptionReason(_opt(row, "reason") or "OTHER"),
            )
            for row in tables["Exceptions"]
        ],
        risk_drivers=[
            RiskDriver(
                id=_text(row, "id"),
                name=_text(row, "name"),
                probability=_float(row, "probability"),
                low=_float(row, "low", default=1.0),
                mode=_float(row, "mode", default=1.2),
                high=_float(row, "high", default=1.5),
                task_ids=[part.strip() for part in _text(row, "task_ids").split(";") if part.strip()],
                owner=_opt(row, "owner") or "",
                group=_opt(row, "group"),
            )
            for row in tables["Risks"]
        ],
        provenance=provenance or Provenance(kind=ProvenanceKind.EXPERIMENT, source=path.name),
    )


def _rows(program: OKRProgram) -> dict[str, list[dict[str, Any]]]:
    return {
        "Program": [
            {
                "id": program.program.id,
                "name": program.program.name,
                "horizon_start": program.program.horizon_start,
                "horizon_end": program.program.horizon_end,
                "status_date": program.program.status_date,
            }
        ],
        "Projects": [
            {
                "id": p.id,
                "code": p.code,
                "name": p.name,
                "priority": p.priority,
                "due_date": p.due_date,
                "deadline": p.deadline,
                "enterprise": p.enterprise,
            }
            for p in program.projects
        ],
        "WBS": [
            {
                "id": n.id,
                "project_id": n.project_id,
                "parent_id": n.parent_id,
                "code": n.code,
                "name": n.name,
                "kind": n.kind.value,
            }
            for n in program.wbs
        ],
        "Tasks": [
            {
                "id": t.id,
                "project_id": t.project_id,
                "wbs_id": t.wbs_id,
                "name": t.name,
                "duration_wd": t.duration_wd,
                "kind": t.kind.value,
                "status": t.status.value,
                "earliest_start": t.earliest_start,
                "latest_finish": t.latest_finish,
                "due_date": t.due_date,
                "deadline": t.deadline,
                "shift_limit_wd": t.shift_limit_wd,
                "pinned": t.pinned,
                "remaining_wd": t.remaining_wd,
                "actual_start": t.actual_start,
                "actual_finish": t.actual_finish,
                "planned_start": t.planned_start,
                "planned_finish": t.planned_finish,
                "okr_stage": t.okr_stage,
            }
            for t in program.tasks
        ],
        "Demands": [
            {"task_id": t.id, "resource_id": d.resource_id, "skill_id": d.skill_id, "units": d.units}
            for t in program.tasks
            for d in t.demands
        ],
        "Dependencies": [
            {
                "src_task_id": e.src_task_id,
                "dst_task_id": e.dst_task_id,
                "type": e.type.value,
                "lag_wd": e.lag_wd,
                "max_lag_wd": e.max_lag_wd,
                "hard": e.hard,
                "source": e.source.value,
            }
            for e in program.dependencies
        ],
        "Resources": [
            {
                "id": r.id,
                "kind": r.kind.value,
                "code": r.code,
                "name": r.name,
                "capacity_units": r.capacity_units,
                "skills": ";".join(r.skills),
                "org_unit": r.org_unit,
            }
            for r in program.resources
        ],
        "Skills": [{"id": s.id, "code": s.code, "name": s.name} for s in program.skills],
        "Exceptions": [
            {
                "resource_id": e.resource_id,
                "start": e.start,
                "end": e.end,
                "units_available": e.units_available,
                "reason": e.reason.value,
            }
            for e in program.capacity_exceptions
        ],
        "Risks": [
            {
                "id": r.id,
                "name": r.name,
                "probability": r.probability,
                "low": r.low,
                "mode": r.mode,
                "high": r.high,
                "task_ids": ";".join(r.task_ids),
                "owner": r.owner,
                "group": r.group,
            }
            for r in program.risk_drivers
        ],
    }


def _table(book: Any, name: str) -> list[dict[str, Any]]:
    if name not in book.sheetnames:
        return []
    sheet = book[name]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    except StopIteration:
        return []
    out = []
    for values in rows:
        record = {header[i]: values[i] if i < len(values) else None for i in range(len(header)) if header[i]}
        marker = record.get(header[0]) if header and header[0] else None
        if marker is None or str(marker).strip() == "" or str(marker).startswith("#"):
            continue
        out.append(record)
    return out


def _raw(row: dict[str, Any], key: str) -> Any:
    value = row.get(key)
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _text(row: dict[str, Any], key: str) -> str:
    value = _raw(row, key)
    if value is None:
        raise ValueError(f"column {key} is empty")
    return str(value)


def _opt(row: dict[str, Any], key: str) -> str | None:
    value = _raw(row, key)
    return None if value is None else str(value)


def _optional_int(row: dict[str, Any], key: str) -> int | None:
    value = _raw(row, key)
    if value is None:
        return None
    return int(value)


def _int(row: dict[str, Any], key: str, default: int = 0) -> int:
    value = _optional_int(row, key)
    return default if value is None else value


def _float(row: dict[str, Any], key: str, default: float | None = None) -> float:
    value = _raw(row, key)
    if value is None:
        if default is None:
            raise ValueError(f"column {key} is empty")
        return default
    return float(str(value).replace(",", "."))


def _bool(row: dict[str, Any], key: str, default: bool = False) -> bool:
    value = _raw(row, key)
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    return str(value).lower() in {"1", "true", "yes", "да"}


def _date(row: dict[str, Any], key: str) -> date | None:
    value = _raw(row, key)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])
